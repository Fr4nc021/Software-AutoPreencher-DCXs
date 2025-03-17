import tkinter as tk
from tkinter import messagebox
from docx import Document

# Lista de modelos disponíveis
modelos_disponiveis = {
    "1": "1.ACERVO.docx",
    "2": "2.DSA.docx",
    "3": "3.IDONEIDADE.docx",
    "4": "4.DECLARAÇÃO DE COMPETIÇÃO.docx",   
}

# Função para preencher o documento Word com os dados
def gerar_documento(dados, modelo_selecionado):
    try:
        modelo_path = modelos_disponiveis.get(modelo_selecionado)  # Obtém o caminho correto do modelo

        if not modelo_path:
            messagebox.showerror("Erro", "Modelo não encontrado!")
            return

        doc = Document(modelo_path)

        # Substituir os placeholders pelos dados inseridos
        for paragrafo in doc.paragraphs:
            for chave, valor in dados.items():
                placeholder = f"{{{{{chave}}}}}"  # Formato {{CHAVE}}
                if placeholder in paragrafo.text:
                    paragrafo.text = paragrafo.text.replace(placeholder, valor)

        # Salvar o documento preenchido
        nome_arquivo = f"{modelo_selecionado}_{dados['NOME COMPLETO'].replace(' ', '_')}.docx"
        doc.save(nome_arquivo)

        messagebox.showinfo("Sucesso", f"Documento gerado: {nome_arquivo}")
    
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao gerar documento: {str(e)}")

# Função para coletar os dados e chamar a função de geração
def coletar_dados():
    dados = {
        "NOME COMPLETO": entry_nome.get(),
        "RG": entry_rg.get(),
        "CPF": entry_cpf.get(),
        "ENDEREÇO COMPLETO": entry_endereco.get(),
        "CIDADE DE NASCIMENTO": entry_cidade.get(),
        "ESTADO DE NASCIMENTO": entry_estado.get(),
        "DATA DE NASCIMENTO": entry_data.get(),
        "NOME PAI": entry_pai.get(),
        "NOME MÃE": entry_mae.get(),
        "DATA DE EXPEDIÇÃO DOCUMENTO": entry_dataexp.get(),
        "ÓRGÃO EXPEDIDOR": entry_orgao.get(),
    }

    # Pega o modelo selecionado no menu suspenso
    modelo_selecionado = modelo_var.get()
    gerar_documento(dados, modelo_selecionado)

# Criar interface gráfica com Tkinter
root = tk.Tk()
root.title("Gerador de Documentos")

# Criando menu suspenso para selecionar o modelo
modelo_var = tk.StringVar(root)
modelo_var.set("1")  # Define um modelo padrão (modelo "1")

label_modelo = tk.Label(root, text="Escolha o modelo:")
label_modelo.grid(row=0, column=0, padx=10, pady=5)
menu_modelo = tk.OptionMenu(root, modelo_var, *modelos_disponiveis.keys())
menu_modelo.grid(row=0, column=1, padx=10, pady=5)

# Criando os campos de entrada
campos = [
    ("Documentos a gerar", "DOCUMENTOS A GERAR"),
    ("Nome Completo", "NOME COMPLETO"),
    ("RG", "RG"),
    ("CPF", "CPF"),
    ("Endereço Completo", "ENDEREÇO COMPLETO"),
    ("Cidade de Nascimento", "CIDADE DE NASCIMENTO"),
    ("Estado de Nascimento", "ESTADO DE NASCIMENTO"),
    ("Data de Nascimento", "DATA DE NASCIMENTO"),
    ("Nome do Pai", "NOME PAI"),
    ("Nome da Mãe", "NOME MÃE"),
    ("Data de Expedição do Documento", "DATA DE EXPEDIÇÃO DOCUMENTO"),
    ("Órgão Expedidor", "ÓRGÃO EXPEDIDOR"),
]

# Dicionário para armazenar os campos de entrada
entradas = {}

# Criar os labels e campos de entrada dinamicamente
for i, (label_texto, chave) in enumerate(campos):
    label = tk.Label(root, text=label_texto + ":")
    label.grid(row=i, column=0, sticky="w", padx=10, pady=5)
    entry = tk.Entry(root)
    entry.grid(row=i, column=1, padx=10, pady=5)
    entradas[chave] = entry

# Criar botões
btn_gerar = tk.Button(root, text="Gerar Documento", command=coletar_dados)
btn_gerar.grid(row=len(campos), column=0, columnspan=2, pady=10)

# Mapear entradas para variáveis globais
entry_nome = entradas["NOME COMPLETO"]
entry_rg = entradas["RG"]
entry_cpf = entradas["CPF"]
entry_endereco = entradas["ENDEREÇO COMPLETO"]
entry_cidade = entradas["CIDADE DE NASCIMENTO"]
entry_estado = entradas["ESTADO DE NASCIMENTO"]
entry_data = entradas["DATA DE NASCIMENTO"]
entry_pai = entradas["NOME PAI"]
entry_mae = entradas["NOME MÃE"]
entry_dataexp = entradas["DATA DE EXPEDIÇÃO DOCUMENTO"]
entry_orgao = entradas["ÓRGÃO EXPEDIDOR"]

# Iniciar interface
root.mainloop()
