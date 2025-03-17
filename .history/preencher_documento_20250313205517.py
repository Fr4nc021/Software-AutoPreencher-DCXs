import tkinter as tk
from tkinter import messagebox
from docx import Document

# Função para preencher o documento Word com os dados
def gerar_documento(dados):
    try:
        modelo_path = "modelo.docx.docx"  # Caminho do modelo do Word
        doc = Document(modelo_path)

        # Substituir os placeholders pelos dados inseridos
        for paragrafo in doc.paragraphs:
            for chave, valor in dados.items():
                placeholder = f"{{{{{chave}}}}}"  # Formato {{CHAVE}}
                if placeholder in paragrafo.text:
                    paragrafo.text = paragrafo.text.replace(placeholder, valor)

        # Salvar o documento preenchido
        nome_arquivo = f"documento_{dados['NOME COMPLETO'].replace(' ', '_')}.docx"
        doc.save(nome_arquivo)

        messagebox.showinfo("Sucesso", f"Documento gerado: {nome_arquivo}")
    
    except Exception as e:
        messagebox.showerror("Erro", f"Erro ao gerar documento: {str(e)}")

# Função para coletar os dados e chamar a função de geração
def coletar_dados():
    dados = {
        "NOME COMPLETO": entry_nome.get(),
        "RG": entry_rg.get(),
        "CPF": entry_cpf.get(),
        "ENDEREÇO COMPLETO": entry_endereco.get(),
        "CIDADE DE NASCIMENTO": entry_cidade.get(),
        "ESTADO DE NASCIMENTO": entry_estado.get(),
        "DATA DE NASCIMENTO": entry_data.get(),
        "NOME PAI": entry_pai.get(),
        "NOME MÃE": entry_mae.get(),
    }

    # Chamar função para gerar documento preenchido
    gerar_documento(dados)

# Criar interface gráfica com Tkinter
root = tk.Tk()
root.title("Preenchimento de Documentos")

# Criando os campos de entrada
campos = [
    ("Nome Completo", "NOME COMPLETO"),
    ("RG", "RG"),
    ("CPF", "CPF"),
    ("Endereço Completo", "ENDEREÇO COMPLETO"),
    ("Cidade de Nascimento", "CIDADE DE NASCIMENTO"),
    ("Estado de Nascimento", "ESTADO DE NASCIMENTO"),
    ("Data de Nascimento", "DATA DE NASCIMENTO"),
    ("Nome do Pai", "NOME PAI"),
    ("Nome da Mãe", "NOME MÃE"),
]

# Dicionário para armazenar os campos de entrada
entradas = {}

# Criar os labels e campos de entrada dinamicamente
for i, (label_texto, chave) in enumerate(campos):
    label = tk.Label(root, text=label_texto + ":")
    label.grid(row=i, column=0, sticky="w", padx=10, pady=5)
    entry = tk.Entry(root)
    entry.grid(row=i, column=1, padx=10, pady=5)
    entradas[chave] = entry

# Criar botões
btn_gerar = tk.Button(root, text="Gerar Documento", command=coletar_dados)
btn_gerar.grid(row=len(campos), column=0, columnspan=2, pady=10)

# Mapear entradas para variáveis globais
entry_nome = entradas["NOME COMPLETO"]
entry_rg = entradas["RG"]
entry_cpf = entradas["CPF"]
entry_endereco = entradas["ENDEREÇO COMPLETO"]
entry_cidade = entradas["CIDADE DE NASCIMENTO"]
entry_estado = entradas["ESTADO DE NASCIMENTO"]
entry_data = entradas["DATA DE NASCIMENTO"]
entry_pai = entradas["NOME PAI"]
entry_mae = entradas["NOME MÃE"]

# Iniciar interface
root.mainloop()
